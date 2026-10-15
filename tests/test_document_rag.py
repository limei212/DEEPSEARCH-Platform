import asyncio
import importlib.util
import tempfile
import threading
import unittest
from pathlib import Path

import numpy as np
from docx import Document
from reportlab.pdfgen import canvas


class FixtureEncoder:
    max_tokens = 40
    def token_count(self, text): return len(text)
    def embed_documents(self, texts):
        return np.array([[text.count('库存')+text.count('inventory'), text.count('品类'), 0.1] for text in texts], dtype=np.float32)
    def embed_query(self, text): return self.embed_documents([text])[0]


class DocumentRAGTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        try:spec=importlib.util.find_spec('app.rag.documents')
        except ModuleNotFoundError:spec=None
        self.assertIsNotNone(spec, '临时文档检索模块尚未实现')
        from app.rag.documents import TaskDocumentIndex, extract_document, split_text
        from app.rag.session import document_scope, current_index
        self.Index=TaskDocumentIndex;self.extract=extract_document;self.split=split_text
        self.scope=document_scope;self.current=current_index
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)

    async def test_text_markdown_docx_table_and_pdf_preserve_locations(self):
        for name in ['a.txt','b.md']:
            (self.root/name).write_text('品类策略应聚焦高复购商品。',encoding='utf-8-sig')
            sections=self.extract(self.root/name)
            self.assertIn('品类', sections[0].text)
            self.assertEqual(sections[0].filename,name)
        doc=Document();doc.add_paragraph('库存管理');table=doc.add_table(rows=1,cols=2)
        table.cell(0,0).text='品类';table.cell(0,1).text='清洁用品';doc.save(self.root/'table.docx')
        sections=self.extract(self.root/'table.docx')
        self.assertTrue(any('清洁用品' in s.text and '表格' in s.location for s in sections))
        pdf=canvas.Canvas(str(self.root/'pages.pdf'));pdf.drawString(20,800,'inventory target 123');pdf.showPage();pdf.drawString(20,800,'second page');pdf.save()
        sections=self.extract(self.root/'pages.pdf')
        self.assertEqual([s.location for s in sections],['第 1 页','第 2 页'])

    async def test_splitting_fits_model_limit_without_losing_tail(self):
        text='库存策略。'*45+'尾部证据'
        chunks=self.split(text,FixtureEncoder())
        self.assertTrue(all(len(chunk)<=40 for chunk in chunks))
        self.assertIn('尾部证据',chunks[-1])

    async def test_query_returns_real_sources_and_releases_only_index(self):
        (self.root/'inventory.txt').write_text('库存库存：目标库存为 123 件。',encoding='utf-8')
        (self.root/'strategy.md').write_text('品类品类：优先高复购商品。',encoding='utf-8')
        async with self.scope(self.root,encoder=FixtureEncoder()) as index:
            hits=await index.search('目标库存',limit=2)
            self.assertEqual(hits['sources'][0]['filename'],'inventory.txt')
            self.assertIn('123',hits['sources'][0]['text'])
            self.assertGreater(index.chunk_count,0)
            self.assertIs(self.current.get(),index)
        self.assertEqual(index.chunk_count,0)
        self.assertIsNone(index.vectors)
        self.assertIsNone(self.current.get())
        self.assertTrue((self.root/'inventory.txt').is_file())
        self.assertFalse(any(p.suffix in ['.db','.sqlite','.npy'] for p in self.root.iterdir()))

    async def test_new_task_rebuilds_and_does_not_include_other_task_documents(self):
        first=self.root/'one';second=self.root/'two';first.mkdir();second.mkdir()
        (first/'first.txt').write_text('库存 123',encoding='utf-8');(second/'second.txt').write_text('库存 456',encoding='utf-8')
        async def query(folder):
            async with self.scope(folder,encoder=FixtureEncoder()) as index:
                return await index.search('库存')
        results=await asyncio.gather(query(first),query(second))
        self.assertEqual([r['sources'][0]['filename'] for r in results],['first.txt','second.txt'])
        (first/'first.txt').write_text('库存 789',encoding='utf-8')
        self.assertIn('789',(await query(first))['sources'][0]['text'])

    async def test_exception_and_cancel_always_release(self):
        (self.root/'a.txt').write_text('库存',encoding='utf-8')
        captured=[]
        async def run(error):
            async with self.scope(self.root,encoder=FixtureEncoder()) as index:
                captured.append(index);await index.search('库存');raise error
        for error in [RuntimeError('失败'),asyncio.CancelledError()]:
            with self.assertRaises(type(error)):await run(error)
            self.assertEqual(captured[-1].chunk_count,0)
            self.assertIsNone(self.current.get())

    async def test_missing_empty_scanned_and_corrupt_documents_report_clear_errors(self):
        async with self.scope(self.root,encoder=FixtureEncoder()) as index:
            with self.assertRaisesRegex(ValueError,'上传'):await index.search('库存')
        (self.root/'empty.txt').write_text('',encoding='utf-8')
        (self.root/'corrupt.pdf').write_bytes(b'not a PDF')
        pdf=canvas.Canvas(str(self.root/'scan.pdf'));pdf.rect(0,0,10,10);pdf.save()
        async with self.scope(self.root,encoder=FixtureEncoder()) as index:
            with self.assertRaisesRegex(ValueError,'文本|解析'):await index.search('库存')

    async def test_selected_filename_cannot_escape_upload_directory(self):
        (self.root/'a.txt').write_text('库存',encoding='utf-8')
        async with self.scope(self.root,encoder=FixtureEncoder()) as index:
            with self.assertRaisesRegex(ValueError,'文件名|路径'):await index.search('库存',filename='../outside.txt')

    async def test_cancel_while_building_cannot_repopulate_released_index(self):
        started=threading.Event();finish=threading.Event()
        class BlockingEncoder(FixtureEncoder):
            def embed_documents(self,texts):
                started.set();finish.wait(5);return super().embed_documents(texts)
        (self.root/'a.txt').write_text('库存',encoding='utf-8');captured=[]
        async def run():
            async with self.scope(self.root,encoder=BlockingEncoder()) as index:
                captured.append(index);await index.search('库存')
        task=asyncio.create_task(run())
        try:
            self.assertTrue(await asyncio.to_thread(started.wait,3))
            task.cancel()
            with self.assertRaises(asyncio.CancelledError):await task
            self.assertTrue(captured[0].closed.is_set())
        finally:finish.set()
        await asyncio.to_thread(lambda:None)
        self.assertEqual(captured[0].chunk_count,0)
        self.assertIsNone(captured[0].vectors)

if __name__=='__main__':unittest.main()
